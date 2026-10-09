"""Build the Sense Foraging Study 2 DRAFT: questionnaire definitions, consent and
debrief HTML, and the SQL that creates the draft study.

A test drive, not the study. It runs on what the platform does today, so it
cannot show the per-participant interleaved order (one sample order is baked in:
seed 12345 from interleave_prototype.py), the background branching (Q7-Q10
always show), the "I do not agree" button, or the return to Prolific. Slugs are
`sfdraft-*` so the real definitions can take clean `sf-*` slugs later.

Sources, read verbatim at build time:
  - pool 4 item wording: receptive_state_item_pool_5.xlsx, column C
  - FFMQ-15, MAIA-2, PHQ-4, BIDR-16: Rev2/Validity_Scales_Appendix.docx
  - consent: Rev2/Consent_Form_v3.docx; debrief: Rev2/Debrief_Form.docx;
    resources: Rev2/Mental_Health_Resources_Sheet.docx

Usage:  python scripts/sense_foraging/build_draft.py   (writes scripts/sense_foraging/draft/)
"""
import html
import json
import re
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

ETHICS = Path(r'I:\My Drive\Ethics\Sense Foraging Questionnaire Zindel\Rev2')
POOL5 = Path(r'I:\Shared drives\SenseForaging\Assessment\MainQuestionnaire\receptive_state_item_pool_5.xlsx')
OUT = Path(__file__).parent / 'draft'

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

# Interleaved order from interleave_prototype.py, seed 12345 (pages 9/8/8/8).
SFS_PAGES = [[30, 23, 17, 14, 12, 25, 2, 24, 10],
             [13, 18, 7, 20, 29, 6, 5, 22],
             [1, 21, 32, 'A', 15, 11, 3, 4],
             [26, 9, 28, 27, 16, 19, 8, 31]]
# Attention checks end in a word, not a digit: the export names a column by an
# item's trailing digits, so `sfs_attn_1` would export as `_1`, i.e. "item 1".
SFS_CHECK = 'sfs_attn_check'

# Interleave rules (pool 5's facet key; verified by interleaveOrder.test.mjs).
def sfs_interleave():
    i = lambda ns: [f'sfs_{n:02d}' for n in ns]
    return {
        'items': i(range(1, 33)) + [SFS_CHECK],
        'page_sizes': [9, 8, 8, 8],
        'clusters': {
            'action': i([1, 6, 7, 8, 9]), 'practice': i([2, 3]), 'drift': i([4, 13, 14]),
            'view_core': i([5, 10, 11, 12]), 'view_stress': i([16, 18]), 'normalizing': i([15, 17]),
            'completeness': i([19, 20, 21]), 'safety': i([22, 23, 24]), 'reward': i([25, 26, 27]),
            'view_openness': i([28, 29, 30, 31, 32]),
        },
        'cluster_min_gap': 3,
        'cluster_max_per_page': 2,
        'groups': [{'items': i([5, 10, 11, 12, 16, 18, 28, 29, 30, 31, 32]), 'min_gap': 2, 'max_per_page': 3}],
        'pairs': [i([4, 13]), i([16, 18]), i([1, 9]), i([6, 16])],
        'pair_min_gap': 3,
        'anchored': {SFS_CHECK: {'min_position': 11, 'max_position': 24, 'not_page_edge': True}},
    }


# One sample scale order for the draft; the real study draws it per participant.
VALIDITY_ORDER = ['maia2', 'phq4', 'bidr16', 'lifesat', 'ffmq15']


# ── docx reading ──────────────────────────────────────────────────────────────
def docx_paragraphs(path):
    """[(style, text)] for every body paragraph, tables flattened row by row."""
    root = ET.fromstring(zipfile.ZipFile(path).read('word/document.xml'))
    out = []

    def para(p):
        text = ''.join(t.text or '' for t in p.iter(W + 't'))
        st = p.find(f'{W}pPr/{W}pStyle')
        style = st.get(W + 'val') if st is not None else ''
        if p.find(f'{W}pPr/{W}numPr') is not None:
            style = style or 'List'
        return style, text

    for el in root.find(W + 'body'):
        if el.tag == W + 'p':
            out.append(para(el))
        elif el.tag == W + 'tbl':
            for p in el.iter(W + 'p'):
                out.append(para(p))
    return out


def section(paras, start_pat, end_pat=None):
    """Paragraphs after the heading matching start_pat, up to end_pat."""
    out, on = [], False
    for style, text in paras:
        if re.match(start_pat, text):
            on = True
            continue
        if on and end_pat and re.match(end_pat, text):
            break
        if on:
            out.append((style, text))
    return out


# ── component helpers ─────────────────────────────────────────────────────────
def scale(points):
    return [{'value': v, 'label': l} for v, l in points]


def likert(cid, question, points):
    return {'id': cid, 'type': 'likert', 'question': question, 'scale': scale(points),
            'required': True, 'allow_pna': True}


def info(cid, title, body):
    return {'id': cid, 'type': 'information', 'title': title, 'body': body,
            'image_url': '', 'image_alt': ''}


def mc(cid, question, options, multiple=False, allow_pna=True):
    opts = []
    for o in options:
        o = dict(o)
        o.setdefault('value', o['id'])
        opts.append(o)
    c = {'id': cid, 'type': 'multiple_choice', 'question': question, 'options': opts,
         'required': True, 'allow_pna': allow_pna}
    if multiple:
        c['allow_multiple'] = True
    return c


def opt(oid, label, **kw):
    return {'id': oid, 'label': label, **kw}


def questionnaire(slug, name, instructions, pages):
    return {'slug': slug, 'name': name, 'questionnaire_type': 'composable',
            'instructions': instructions, 'pages': pages}


def chunk(seq, sizes):
    out, i = [], 0
    for n in sizes:
        out.append(seq[i:i + n])
        i += n
    assert i == len(seq), (i, len(seq))
    return out


# ── instruments ───────────────────────────────────────────────────────────────
SFS_POINTS = [(1, 'Strongly disagree'), (2, 'Disagree'), (3, 'Mildly disagree'),
              (4, 'Mildly agree'), (5, 'Agree'), (6, 'Strongly agree')]


def build_sfs(slug='sfdraft-sfs-pool4', name='Sense Foraging Scale (DRAFT, pool 4 wording)', interleave=False):
    wb = openpyxl.load_workbook(POOL5)
    rows = {r[0].value: r[2].value for r in wb['Item Pool'].iter_rows(min_row=2) if r[0].value}
    definition = next(r[1].value for r in wb['Instructions'].iter_rows()
                      if r[0].value == 'Pool 4 definition shown before the items')
    assert len(rows) == 32
    check = likert(SFS_CHECK, 'To show that you are reading each statement, please select '
                   'Disagree for this one.', SFS_POINTS)
    pages = []
    for p, items in enumerate(SFS_PAGES, 1):
        comps = [check if i == 'A' else likert(f'sfs_{i:02d}', rows[i], SFS_POINTS) for i in items]
        pages.append({'id': f'sfs_p{p}', 'components': comps})
    instructions = (definition + '\n\nPlease rate how much you feel that each statement describes '
                    'you right now, on a scale of 1-6.')
    q = questionnaire(slug, name, instructions, pages)
    if interleave:
        # The written pages (seed 12345, itself a valid order) stay as the
        # fallback the renderer uses if no order can be drawn.
        q['interleave'] = sfs_interleave()
    return q


def appendix_items(paras, heading):
    sec = section(paras, heading, r'^#?\s*\d\.\s|^References')
    return [re.sub(r'\s*\(R\)\s*$', '', t).strip() for s, t in sec if s and t.strip()
            and not t.startswith(('Citation', 'Licensing', 'Instructions', 'Scoring', 'Note'))]


def build_validity():
    paras = docx_paragraphs(ETHICS / 'Validity_Scales_Appendix.docx')
    texts = [t for _, t in paras]

    def items_after(title_start, n):
        i = next(k for k, t in enumerate(texts) if t.startswith(title_start))
        found = []
        for style, t in paras[i + 1:]:
            if t.startswith('Scoring'):
                break
            if style and t.strip() and not t.startswith(('Citation', 'Licensing', 'Instructions')):
                found.append(re.sub(r'\s*\(R\)\s*$', '', t).strip())
        assert len(found) == n, (title_start, len(found))
        return found

    ffmq = items_after('1. Five Facet', 15)
    maia = items_after('2. Multidimensional', 37)
    phq = items_after('4. Patient Health', 4)
    # The SDE and IM subscales are one paragraph each, items numbered inline.
    bidr_par = ' '.join(t for t in texts if t.startswith(('Self-Deceptive Enhancement',
                                                           'Impression Management')))
    bidr = [re.sub(r'\s*\(R\)\s*$', '', m).strip()
            for m in re.split(r'\s(?:\d{1,2})\.\s', ' ' + re.sub(
                r'Impression Management \(IM\)', '', bidr_par.replace('Self-Deceptive Enhancement (SDE)', '')))
            if m.strip()]
    assert len(bidr) == 16, len(bidr)

    out = {}
    pts = [(1, 'Never or very rarely true'), (2, 'Rarely true'), (3, 'Sometimes true'),
           (4, 'Often true'), (5, 'Very often or always true')]
    comps = [likert(f'ffmq_{i:02d}', t, pts) for i, t in enumerate(ffmq, 1)]
    out['ffmq15'] = questionnaire(
        'sfdraft-ffmq15', 'FFMQ-15 (DRAFT)',
        'Please use the 1 (never or very rarely true) to 5 (very often or always true) scale provided '
        'to indicate how true the below statements are of you. Please respond according to what '
        'really reflects your experience rather than what you think your experience should be.',
        [{'id': f'ffmq_p{k}', 'components': c} for k, c in enumerate(chunk(comps, [8, 7]), 1)])

    pts = [(0, '0 Never'), (1, '1'), (2, '2'), (3, '3'), (4, '4'), (5, '5 Always')]
    comps = [likert(f'maia_{i:02d}', t, pts) for i, t in enumerate(maia, 1)]
    comps.insert(19, likert('maia_attn_check', 'To show that you are reading each statement, please '
                            'select 2 for this one.', pts))
    out['maia2'] = questionnaire(
        'sfdraft-maia2', 'MAIA-2 (DRAFT)',
        'Below you will find a list of statements. Please indicate how often each statement applies '
        'to you generally in daily life (0 = Never, 5 = Always).',
        [{'id': f'maia_p{k}', 'components': c} for k, c in enumerate(chunk(comps, [8, 8, 8, 7, 7]), 1)])

    pts = [(0, 'Not at all'), (1, 'Several days'), (2, 'More than half the days'), (3, 'Nearly every day')]
    out['phq4'] = questionnaire(
        'sfdraft-phq4', 'PHQ-4 (DRAFT)',
        'Over the last two weeks, how often have you been bothered by the following problems?',
        [{'id': 'phq4_p1', 'components': [likert(f'phq4_{i}', t, pts) for i, t in enumerate(phq, 1)]}])

    pts = [(1, '1 Not true'), (2, '2'), (3, '3'), (4, '4 Somewhat true'), (5, '5'), (6, '6'), (7, '7 Very true')]
    comps = [likert(f'bidr_{i:02d}', t, pts) for i, t in enumerate(bidr, 1)]
    out['bidr16'] = questionnaire(
        'sfdraft-bidr16', 'BIDR-16 (DRAFT)',
        'Using the scale below, please indicate how much you agree with each statement. '
        '(1 = Not True, 4 = Somewhat True, 7 = Very True)',
        [{'id': f'bidr_p{k}', 'components': c} for k, c in enumerate(chunk(comps, [8, 8]), 1)])
    return out


def build_background(slug='sfdraft-background', name='Background Questionnaire (DRAFT)', branching=False):
    yn = [opt('yes', 'Yes'), opt('no', 'No')]
    page_a = [
        mc('bg_age', 'What is your age?',
           [opt('age', 'My age in years', response_type='number', min=18, max=100),
            opt('prefer_not', 'Prefer not to answer')], allow_pna=False),
        mc('bg_gender', 'What is your gender?',
           [opt('man', 'Man'), opt('woman', 'Woman'), opt('non_binary', 'Non-binary'),
            opt('self_describe', 'Prefer to self-describe', response_type='text'),
            opt('prefer_not', 'Prefer not to say')], allow_pna=False),
        mc('bg_race', 'What is your race/ethnicity? Select all that apply.',
           [opt('east_asian', 'East Asian'), opt('south_asian', 'South Asian'),
            opt('southeast_asian', 'Southeast Asian'), opt('black', 'Black'),
            opt('hispanic_latin', 'Hispanic or Latino/a/x'),
            opt('indigenous', 'Indigenous (e.g., First Nations, Native American, Aboriginal and '
                'Torres Strait Islander, Māori)'),
            opt('mena', 'Middle Eastern or North African'), opt('pacific', 'Pacific Islander'),
            opt('white', 'White'), opt('mixed', 'Mixed or multiple'),
            opt('other', 'Other (please specify)', response_type='text'),
            opt('prefer_not', 'Prefer not to say', exclusive=True)], multiple=True, allow_pna=False),
        mc('bg_country', 'What is your current country of residence?',
           [opt('us', 'United States'), opt('uk', 'United Kingdom'), opt('ca', 'Canada'),
            opt('au', 'Australia'), opt('ie', 'Ireland'), opt('nz', 'New Zealand'),
            opt('other', 'Other (please specify)', response_type='text')]),
        mc('bg_education', 'What is the highest level of education you have completed?',
           [opt('lt_hs', 'Less than high school'), opt('hs', 'High school or equivalent'),
            opt('some_college', 'Some college'), opt('bachelors', 'Bachelor’s degree'),
            opt('graduate', 'Graduate or professional degree'),
            opt('prefer_not', 'Prefer not to say')], allow_pna=False),
    ]
    page_b = [
        mc('bg_practice_regular', 'Have you ever practiced meditation, yoga, or another contemplative '
           'or mind-body practice regularly (i.e., at least once a week for a period of two months or '
           'more)?', yn),
    ]
    page_c = [
        mc('bg_practice_types', 'Which practice(s) have you engaged in regularly? Select all that apply.',
           [opt('meditation', 'Meditation (e.g., mindfulness, focused attention, loving-kindness)'),
            opt('yoga', 'Yoga'), opt('taichi_qigong', 'Tai chi / Qigong'),
            opt('contemplative_prayer', 'Contemplative prayer or religious contemplative practice'),
            opt('martial_arts', 'Martial arts with a contemplative/meditative component'),
            opt('other', 'Other', response_type='text')], multiple=True),
        mc('bg_practice_years', 'In total, for how many years have you practiced regularly (across all '
           'practices listed above)?',
           [opt('years', 'Years', response_type='number', min=0, max=90),
            opt('prefer_not', 'Prefer not to answer')], allow_pna=False),
        mc('bg_practice_freq', 'Over the past 3 months, how often have you practiced?',
           [opt('none', 'Not at all'), opt('lt_weekly', 'Less than once a week'),
            opt('1_2_week', '1–2 times a week'), opt('3_5_week', '3–5 times a week'),
            opt('daily', 'Daily or near-daily')]),
        mc('bg_practice_training', 'Have you ever completed a formal, teacher-led training or program in a '
           'contemplative practice (e.g., an MBSR course, a yoga teacher training, a meditation '
           'retreat)?', yn),
    ]
    if branching:
        # Q7-Q10 only after "Yes" to Q6 (approved Background Questionnaire).
        for c in page_c:
            c['show_if'] = {'component': 'bg_practice_regular', 'equals': 'yes'}
    else:
        page_c.insert(0, info('bg_draft_branch_note', 'Draft note',
            'In the final study this page appears only if you answered Yes on the previous page. '
            'The draft cannot skip it yet: if you answered No, choose "Prefer not to answer" here.'))
    rel = [(1, '1 Not at all religious'), (2, '2'), (3, '3'), (4, '4'), (5, '5 Very religious')]
    spi = [(1, '1 Not at all spiritual'), (2, '2'), (3, '3'), (4, '4'), (5, '5 Very spiritual')]
    page_d = [
        likert('bg_religious', 'To what extent do you consider yourself a religious person?', rel),
        likert('bg_spiritual', 'To what extent do you consider yourself a spiritual person?', spi),
        mc('bg_tradition', 'Do you currently identify with a religious or spiritual tradition? Select the '
           'one that best applies.',
           [opt('none', 'No religious/spiritual affiliation'), opt('christian', 'Christian'),
            opt('muslim', 'Muslim'), opt('jewish', 'Jewish'), opt('buddhist', 'Buddhist'),
            opt('hindu', 'Hindu'), opt('sikh', 'Sikh'),
            opt('sbnr', 'Spiritual but not religious (unaffiliated)'),
            opt('other', 'Other', response_type='text'), opt('prefer_not', 'Prefer not to say')],
           allow_pna=False),
    ]
    return questionnaire(
        slug, name,
        'A few questions about you, including any regular meditation, yoga, or other mind-body practice. '
        'None of your answers affect whether you take part.',
        [{'id': 'bg_about', 'components': page_a}, {'id': 'bg_practice', 'components': page_b},
         {'id': 'bg_practice_detail', 'components': page_c}, {'id': 'bg_spirit', 'components': page_d}])


# ── consent and debrief HTML ──────────────────────────────────────────────────
DRAFT_BANNER = ('<p style="border:2px solid #c00000;border-radius:12px;padding:8px 16px;color:#c00000">'
                '<strong>DRAFT for review — not a live study.</strong> No one is recruited to this page.'
                '</p>')


def to_html(paras, skip=lambda t: False, h_level=2):
    out, in_list = [], False
    for style, text in paras:
        text = text.strip()
        if not text or skip(text):
            continue
        is_list = style == 'List' or style.lower().startswith('list')
        if in_list and not is_list:
            out.append('</ul>')
            in_list = False
        esc = html.escape(text)
        if style.lower().startswith('heading'):
            out.append(f'<h{h_level + 1 if style.endswith("3") else h_level}>{esc}</h{h_level + 1 if style.endswith("3") else h_level}>')
        elif is_list:
            if not in_list:
                out.append('<ul>')
                in_list = True
            out.append(f'<li>{esc}</li>')
        else:
            out.append(f'<p>{esc}</p>')
    if in_list:
        out.append('</ul>')
    return '\n'.join(out)


def build_consent():
    paras = docx_paragraphs(ETHICS / 'Consent_Form_v3.docx')
    subheads = {'Data Retention and Disposition', 'Research Ethics Board Access',
                'Future Use of Your Data (Optional)', 'Main Study Consent (required)',
                'Repository / Future Use Consent (optional)'}
    paras = [('Heading3', t) if t.strip() in subheads else (s, t) for s, t in paras]
    # The checkbox lines are rendered by ConsentGate itself (agree checkbox and the
    # separate Borealis Yes/No), so the form's own [ ] lines are left out.
    skip = lambda t: t.startswith('[') or t.startswith('REVISED')
    return DRAFT_BANNER + '\n' + to_html(paras, skip)


def build_debrief():
    deb = docx_paragraphs(ETHICS / 'Debrief_Form.docx')
    cut = next(i for i, (_, t) in enumerate(deb) if t.startswith('Mental Health and Wellbeing Resources'))
    closing = [p for p in deb[cut:] if p[1].startswith('Thank you again')]
    res = docx_paragraphs(ETHICS / 'Mental_Health_Resources_Sheet.docx')
    start = next(i for i, (_, t) in enumerate(res) if t.startswith('Taking part'))
    resources = [('Heading2', 'Mental Health and Wellbeing Resources')] + res[start:]
    return (DRAFT_BANNER + '\n' + to_html(deb[:cut]) + '\n' + to_html(resources) + '\n'
            + to_html(closing))


# ── SQL ───────────────────────────────────────────────────────────────────────
def lit(s):
    return "'" + s.replace("'", "''") + "'"


def main():
    OUT.mkdir(exist_ok=True)
    defs = {'background': build_background(), 'sfs': build_sfs(), **build_validity()}
    for key, d in defs.items():
        (OUT / f"{d['slug']}.json").write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding='utf-8')
    consent, debrief = build_consent(), build_debrief()
    (OUT / 'consent.html').write_text(consent, encoding='utf-8')
    (OUT / 'debrief.html').write_text(debrief, encoding='utf-8')

    ids = {k: str(uuid.uuid4()) for k in ['study', 'consent', 'debrief', 'template', 'session']}
    q_ids = {k: str(uuid.uuid4()) for k in defs}
    steps = ['background', 'sfs'] + VALIDITY_ORDER
    labels = {'background': 'Background questionnaire', 'sfs': 'Sense Foraging Scale',
              'ffmq15': 'FFMQ-15', 'maia2': 'MAIA-2', 'phq4': 'PHQ-4', 'bidr16': 'BIDR-16',
              'lifesat': 'Life satisfaction'}
    sql = ['BEGIN;']
    for k, d in defs.items():
        sql.append(f"INSERT INTO questionnaires (id, slug, name, definition) VALUES ({lit(q_ids[k])}, "
                   f"{lit(d['slug'])}, {lit(d['name'])}, {lit(json.dumps(d, ensure_ascii=False))}::jsonb);")
    sql.append(
        "INSERT INTO studies (id, name, public_title, delivery_mode, active, consent_required, "
        "offer_repository_consent, allow_external_enrollment, external_enrollment_source, "
        "reminders_enabled, compensation_kind, reply_to_email) VALUES ("
        f"{lit(ids['study'])}, 'Sense Foraging Study 2 (DRAFT, pool 4 wording)', "
        "'Mindfulness, Body Awareness & Everyday Experience Survey', 'online_single', true, true, "
        "true, true, 'prolific', false, 'pay', 'norman.farb@utoronto.ca');")
    sql.append(f"INSERT INTO study_consent_forms (id, study_id, html_content) VALUES ({lit(ids['consent'])}, "
               f"{lit(ids['study'])}, {lit(consent)});")
    sql.append(f"INSERT INTO study_debrief_forms (id, study_id, html_content) VALUES ({lit(ids['debrief'])}, "
               f"{lit(ids['study'])}, {lit(debrief)});")
    sql.append(f"UPDATE studies SET active_consent_form_id = {lit(ids['consent'])}, "
               f"active_debrief_form_id = {lit(ids['debrief'])} WHERE id = {lit(ids['study'])};")
    sql.append(f"INSERT INTO session_templates (id, label, folder, description) VALUES ({lit(ids['template'])}, "
               "'SenseForaging_Study2_DRAFT', 'Sense Foraging', 'Draft test drive of Study 2: background, the 32 "
               "items in one sample interleaved order (seed 12345), then the validity scales in one sample "
               "order. Built by scripts/sense_foraging/build_draft.py.');")
    order = 0
    for k in steps:
        if k == 'lifesat':
            sql.append("INSERT INTO session_template_nodes (session_template_id, order_index, activity_id, label) "
                       f"VALUES ({lit(ids['template'])}, {order}, 'e6ff9b25-d826-4561-8121-1201798948dc', 'Life satisfaction');")
        else:
            sql.append("INSERT INTO session_template_nodes (session_template_id, order_index, questionnaire_id, label) "
                       f"VALUES ({lit(ids['template'])}, {order}, {lit(q_ids[k])}, {lit(labels[k])});")
        order += 1
    sql.append("INSERT INTO session_template_nodes (session_template_id, order_index, activity_id, label) "
               f"VALUES ({lit(ids['template'])}, {order}, 'f79f90f7-9f59-44a6-8cf7-4ca45a76298a', 'Debrief');")
    # send_time must be set: auto-enroll's single-shot path copies it into
    # participant_schedule.send_time, which is NOT NULL. Left null, every join
    # fails after the enrollment is created (found on the first test, 2026-10-09).
    sql.append("INSERT INTO study_sessions (id, study_id, session_template_id, day_number, send_time, label, "
               f"order_index, node_key, link_expires_hours) VALUES ({lit(ids['session'])}, {lit(ids['study'])}, "
               f"{lit(ids['template'])}, 1, '09:00', 'Single session', 0, 'sf_single', 48);")
    sql.append('COMMIT;')
    (OUT / 'draft.sql').write_text('\n'.join(sql), encoding='utf-8')
    (OUT / 'ids.json').write_text(json.dumps({**ids, 'questionnaires': q_ids}, indent=2), encoding='utf-8')
    for k, d in defs.items():
        n = sum(1 for p in d['pages'] for c in p['components'] if c['type'] != 'information')
        print(f"{d['slug']:24s} pages={len(d['pages'])} answerable={n}")
    print('study', ids['study'])

    # Study 1 (EFA): the real definitions. They update the existing sf-* rows in
    # place; no participant has answered them (test enrollments only).
    s1 = OUT.parent / 'study1'
    s1.mkdir(exist_ok=True)
    study1 = [build_background('sf-background', 'About you', branching=True),
              build_sfs('sf-sfs', 'How you pay attention', interleave=True)]
    upd = ['BEGIN;']
    for d in study1:
        (s1 / f"{d['slug']}.json").write_text(json.dumps(d, ensure_ascii=False, indent=2), encoding='utf-8')
        upd.append(f"UPDATE questionnaires SET definition = {lit(json.dumps(d, ensure_ascii=False))}::jsonb, "
                   f"name = {lit(d['name'])}, updated_at = now() WHERE slug = {lit(d['slug'])};")
    upd.append('COMMIT;')
    (s1 / 'study1_update.sql').write_text('\n'.join(upd), encoding='utf-8')
    print('study 1 definitions:', ', '.join(d['slug'] for d in study1))


if __name__ == '__main__':
    main()
