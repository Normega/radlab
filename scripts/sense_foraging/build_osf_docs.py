"""OSF files for Sense Foraging Study 1 (exploratory factor analysis of item pool 6).

Writes to the ethics project folder, OSF_Study1/:
  - Study1_OSF_Preregistration.docx: docs/markdowns/sense_foraging_study1_preregistration.md
    rendered to Word (its headings follow OSF's standard preregistration template)
  - Study1_Supplementary_Materials.docx: the materials exactly as administered, from
    the live definitions (study1/sf-background.json, study1/sf-pool6.json, both
    checked equal to the database): background questions with options and
    branching, the scale instructions, response format, item order rules, and
    every item with its id, keying, hypothesized factor, facet and partner
  - Study1_item_key.csv: the same item key, for analysis

Usage:  python scripts/sense_foraging/build_osf_docs.py
"""
import csv
import json
import re
import sys
from pathlib import Path

import docx
from docx.shared import Pt

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from build_pool6 import POOL5, SLOTS  # noqa: E402

ROOT = HERE.parents[1]
OUT = Path(r'I:\My Drive\Ethics\Sense Foraging Questionnaire Zindel\OSF_Study1')
PREREG = ROOT / 'docs' / 'markdowns' / 'sense_foraging_study1_preregistration.md'
BG = json.loads((HERE / 'study1' / 'sf-background.json').read_text(encoding='utf-8'))
SF = json.loads((HERE / 'study1' / 'sf-pool6.json').read_text(encoding='utf-8'))


def styled(d):
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    return d


def add_runs(p, text):
    """Markdown **bold**, *italic* and `code` inside a paragraph."""
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text):
        if not part:
            continue
        if part.startswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('*'):
            p.add_run(part[1:-1]).italic = True
        elif part.startswith('`'):
            r = p.add_run(part[1:-1])
            r.font.name = 'Consolas'
        else:
            p.add_run(part)


def markdown_to_docx(md, path):
    d = styled(docx.Document())
    para = []

    def flush():
        if para:
            add_runs(d.add_paragraph(), ' '.join(para))
            para.clear()

    for line in md.splitlines():
        s = line.rstrip()
        m = re.match(r'^(#{1,4})\s+(.*)', s)
        lm = re.match(r'^(\s*)(?:[-*]|(\d+)\.)\s+(.*)', s)
        if m:
            flush()
            d.add_heading(m.group(2), level=len(m.group(1)))
        elif lm:
            flush()
            nested = len(lm.group(1)) >= 2
            style = ('List Number' if lm.group(2) else 'List Bullet') + (' 2' if nested else '')
            add_runs(d.add_paragraph(style=style), lm.group(3))
        elif not s:
            flush()
        else:
            para.append(s.strip())
    flush()
    d.save(path)


def item_key():
    comps = {c['id']: c for p in SF['pages'] for c in p['components']}
    factor_of = {i: f for f, ids in SF['factors'].items() for i in ids}
    facet = {i['n']: i['facet'] for i in POOL5['items']}
    partner = {}
    for k, (slot, f, a, b, pick) in enumerate(SLOTS):
        n = 35 + k
        facet[n] = f
        partner[n] = (a if pick == 'a' else b)[0]
        partner[partner[n]] = n
    rev = set(SF['reverse_keyed'])
    rows = []
    for n in range(1, 49):
        cid = f'sf6_{n}'
        rows.append({'item': n, 'id': cid, 'wording': comps[cid]['question'],
                     'keying': 'reverse' if cid in rev else 'positive', 'hypothesized_factor': factor_of[cid],
                     'pool5_facet': facet[n], 'paired_with': partner.get(n, '')})
    for cid in ('sf6_attn_disagree', 'sf6_attn_agree'):
        rows.append({'item': '', 'id': cid, 'wording': comps[cid]['question'], 'keying': 'attention check',
                     'hypothesized_factor': '', 'pool5_facet': '',
                     'paired_with': 'correct = 2' if 'disagree' in cid else 'correct = 5'})
    return rows


def option_text(c):
    if c['type'] == 'likert':
        return ' / '.join(f"{s['value']}" + (f" {s['label']}" if s.get('label') and s['label'] != str(s['value']) else '')
                          for s in c['scale'])
    parts = []
    for o in c.get('options', []):
        t = o['label']
        if o.get('response_type') == 'text':
            t += ' [free text]'
        elif o.get('response_type') == 'number':
            t += ' [number]'
        if o.get('exclusive'):
            t += ' [exclusive]'
        parts.append(t)
    return ' / '.join(parts)


def supplementary(rows):
    d = styled(docx.Document())
    d.add_heading('Sense Foraging Scale, Study 1: materials as administered', level=1)
    d.add_paragraph('Supplementary file to the OSF preregistration. Generated from the study definitions that ran '
                    'on the RADlab platform (questionnaires "sf-background" and "sf-pool6"), checked identical to '
                    'the live database. University of Toronto REB Protocol #00051180.')
    d.add_paragraph('Session order: consent (including the optional data-repository question) → background '
                    'questions → Sense Foraging items → debrief with mental-health resources → return to Prolific. '
                    'The consent and debrief text are filed with the REB amendment (Rev3).')

    d.add_heading('A. Background questions', level=2)
    d.add_paragraph(f'Instructions shown: "{BG["instructions"]}"')
    d.add_paragraph('Every question is required, and each offers "Prefer not to answer" (or "Prefer not to say"). '
                    'Questions shown only after a "Yes" to regular practice are marked; when hidden, they are '
                    'recorded as "not applicable".')
    t = d.add_table(rows=1, cols=4)
    t.style = 'Table Grid'
    for cell, h in zip(t.rows[0].cells, ('Page', 'Question (id)', 'Response options', 'Format')):
        cell.text = h
    for pi, page in enumerate(BG['pages'], 1):
        for c in page['components']:
            if c['type'] == 'information':
                continue
            fmt = []
            fmt.append('select all that apply' if c.get('allow_multiple') else
                       ('rating' if c['type'] == 'likert' else 'select one'))
            if c.get('allow_pna'):
                fmt.append('+ Prefer not to answer')
            if c.get('show_if'):
                fmt.append(f"shown only if {c['show_if']['component']} = {c['show_if'].get('equals')}")
            r = t.add_row().cells
            r[0].text, r[1].text, r[2].text, r[3].text = str(pi), f"{c['question']} ({c['id']})", option_text(c), '; '.join(fmt)

    d.add_heading('B. Sense Foraging items (item pool 6)', level=2)
    d.add_paragraph(f'Instructions shown before the items: "{SF["instructions"]}"')
    comps = {c['id']: c for p in SF['pages'] for c in p['components']}
    d.add_paragraph('Response scale (every point labelled, no midpoint): ' + option_text(comps['sf6_1'])
                    + '. A small "Prefer not to answer" is offered under every item except the attention checks.')
    spec = SF['interleave']
    d.add_paragraph('Presentation: 50 items (48 + 2 attention checks) on 6 pages of '
                    + '/'.join(map(str, spec['page_sizes'])) + ' items, each item a card with its own labelled '
                    'options. Order: drawn per participant, seeded so a reload shows the same order, at random from '
                    'all orders satisfying the rules below. Each item\'s position is recorded.')
    for rule in [
        f"Items of the same hypothesized factor at least {spec['cluster_min_gap']} positions apart, at most "
        f"{spec['cluster_max_per_page']} per page.",
        'Near-duplicate pairs at least ' + str(spec['pair_min_gap']) + ' positions apart: '
        + ', '.join(f"{a[4:]}–{b[4:]}" for a, b in spec['pairs']) + '.',
        'Each reverse-keyed item on a different page from its partner.',
        'Reverse-keyed items: 2–3 per page, never adjacent to each other, never in position 1, never beside an '
        'attention check; no run of more than 5 positively keyed items.',
        'Attention check "select Disagree" in positions 3–24, "select Agree" in positions 29–48; never first or last '
        'on a page.',
    ]:
        d.add_paragraph(rule, style='List Bullet')

    t = d.add_table(rows=1, cols=6)
    t.style = 'Table Grid'
    for cell, h in zip(t.rows[0].cells, ('#', 'Wording', 'Keying', 'Hypothesized factor', 'Pool 5 facet', 'Pair')):
        cell.text = h
    for row in rows:
        r = t.add_row().cells
        r[0].text = str(row['item'] or row['id'])
        r[1].text, r[2].text, r[3].text = row['wording'], row['keying'], row['hypothesized_factor']
        r[4].text, r[5].text = row['pool5_facet'], str(row['paired_with'])
    for table in d.tables:
        for tr in table.rows:
            for cell in tr.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
        for cell in table.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True

    d.add_heading('C. Scoring notes', level=2)
    for line in [
        'Items are stored as 1–6 exactly as answered; "Prefer not to answer" is stored as "pna", and hidden '
        'background questions as "not_applicable".',
        'Reverse-keyed items (35–48) are recoded as 7 − x for scoring; they are left unrecoded in models with a '
        'random-intercept acquiescence factor.',
        'Each reverse-keyed item was written as the logical opposite of one positive item (column "Pair"), without '
        'negations, so that agreeing with both signals inconsistency or a general tendency to agree.',
        'Hypothesized factors group pool 5\'s 11 facets so every factor has at least 3 positively keyed items: '
        'View includes View under stress; Mode awareness = Practice + Drift awareness; Noting without elaborating '
        '= Completeness + Labelling; Stress and safety = Normalising stress + Safety.',
    ]:
        d.add_paragraph(line, style='List Bullet')
    d.save(OUT / 'Study1_Supplementary_Materials.docx')


def main():
    OUT.mkdir(exist_ok=True)
    markdown_to_docx(PREREG.read_text(encoding='utf-8'), OUT / 'Study1_OSF_Preregistration.docx')
    rows = item_key()
    with open(OUT / 'Study1_item_key.csv', 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    supplementary(rows)
    print('wrote', OUT)


if __name__ == '__main__':
    main()
