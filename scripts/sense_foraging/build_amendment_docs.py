"""Word documents for the Study 1 minor amendment (REB #00051180), October 2026.

Writes, beside the approved Rev2 package, a folder Rev3_Study1_Amendment with a
clean and a highlighted copy of each changed document:
  - Consent_Form_v4_Study1          (Rev2/Consent_Form_v3.docx, five paragraphs replaced)
  - Debrief_Form_Study1             (Rev2/Debrief_Form.docx, three paragraphs replaced)
  - Mental_Health_Resources_Sheet_Study1 (one paragraph replaced)
  - Background_Questionnaire_Study1 (Rev2/Background_Questionnaire.docx, response options specified)
  - Sense_Foraging_Questionnaire_Study1 (new: the 48 statements and two checks as shown,
    and a table of every change from the approved v1.2 questionnaire)

The consent, debrief and resources edits are the same ones the live Study 1 forms
carry (build_study1_forms.py), applied to the approved Word files paragraph by
paragraph; the build fails if an approved paragraph is not found verbatim. The
questionnaire wording is read from study1/sf-pool6.json, the live definition.

Usage:  python scripts/sense_foraging/build_amendment_docs.py
"""
import json
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from build_study1_forms import CONSENT, DEBRIEF  # noqa: E402

ETHICS = Path(r'I:\My Drive\Ethics\Sense Foraging Questionnaire Zindel')
REV2 = ETHICS / 'Rev2'
OUT = ETHICS / 'Rev3_Study1_Amendment'
DEF = json.loads((HERE / 'study1' / 'sf-pool6.json').read_text(encoding='utf-8'))
AMENDMENT = 'Study 1 amendment, October 2026'


def set_text(p, text, mark):
    """Replace a paragraph's text, keeping its first run's formatting."""
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if mark:
        runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW


def edit_doc(src, dst_stem, edits, mark_note=None, unbold_body=False):
    """edits: [(old paragraph text, new text)]. Writes a clean and a highlighted copy."""
    for mark in (False, True):
        d = docx.Document(src)
        if unbold_body:
            # Rev2 marked its own changes (the Borealis consent split) in bold body
            # text; carried over, that bold would read as part of this amendment.
            for p in d.paragraphs:
                if len(p.text) > 80 and not p.text.startswith('['):
                    for r in p.runs:
                        r.bold = None
        for old, new in edits:
            hits = [p for p in d.paragraphs if p.text.strip() == old.strip()]
            assert len(hits) == 1, f'{src.name}: approved paragraph not found once: {old[:70]}'
            set_text(hits[0], new, mark)
        if mark and mark_note:
            first = d.paragraphs[0]
            note = first.insert_paragraph_before(mark_note)
            note.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
            note.runs[0].font.bold = True
        name = f'{dst_stem}{"_changes_highlighted" if mark else ""}.docx'
        d.save(OUT / name)
        print('wrote', name)


NOTE = f'{AMENDMENT}: changed text is highlighted. All other text is unchanged from the approved Rev2 version.'


def consent():
    edits = [(old, new) for _, old, new in CONSENT]
    d0 = docx.Document(REV2 / 'Consent_Form_v3.docx')
    revised = next(p.text for p in d0.paragraphs if p.text.startswith('REVISED'))
    edits.append((revised, f'REVISED ({AMENDMENT}): Study 1 administers the background questions and the '
                           'Sense Foraging questionnaire only (48 statements and two attention checks, about '
                           '10 minutes). The validity measures move to Study 2.'))
    edit_doc(REV2 / 'Consent_Form_v3.docx', 'Consent_Form_v4_Study1', edits, NOTE, unbold_body=True)


def debrief():
    by_where = {w: (o, n) for w, o, n in DEBRIEF}
    edits = [by_where['About This Study'], by_where['About This Study (second paragraph)'],
             by_where['Your Data (second paragraph)']]
    edit_doc(REV2 / 'Debrief_Form.docx', 'Debrief_Form_Study1', edits, NOTE)
    edit_doc(REV2 / 'Mental_Health_Resources_Sheet.docx', 'Mental_Health_Resources_Sheet_Study1',
             [by_where['Mental Health and Wellbeing Resources (introduction)']], NOTE)


def background():
    bg = json.loads((HERE / 'study1' / 'sf-background.json').read_text(encoding='utf-8'))
    comp = {c['id']: c for p in bg['pages'] for c in p['components']}
    labels = lambda cid: ' / '.join(o['label'] for o in comp[cid]['options'])
    src = REV2 / 'Background_Questionnaire.docx'
    d0 = docx.Document(src)
    find = lambda start: next(p.text for p in d0.paragraphs if p.text.startswith(start))
    edits = [
        (find('Administered after consent'),
         'Administered after consent, before the Sense Foraging questionnaire. No responses here affect '
         'eligibility — everyone who reaches this point completes the full study. Every question offers '
         '"Prefer not to answer" (or "Prefer not to say").'),
        (find('What is your age?'), 'What is your age? (numeric entry, or "Prefer not to answer")'),
        (find('What is your race/ethnicity?'),
         f'What is your race/ethnicity? (checklist, select all that apply: {labels("bg_race")})'),
        (find('What is your current country of residence?'),
         f'What is your current country of residence? ({labels("bg_country")} / Prefer not to answer)'),
        (find('In total, for how many years'),
         'In total, for how many years have you practiced regularly (across all practices listed above)? '
         '(numeric entry, years, or "Prefer not to answer")'),
    ]
    edit_doc(src, 'Background_Questionnaire_Study1', edits, NOTE)


def questionnaire():
    comps = {c['id']: c for p in DEF['pages'] for c in p['components']}
    reversed_ids = set(DEF['reverse_keyed'])
    factor_of = {i: f for f, ids in DEF['factors'].items() for i in ids}
    v12 = docx.Document(ETHICS / 'Sense Foraging and Other Scales v1.2.docx')
    approved = {}
    for p in v12.paragraphs:
        m = re.match(r'^(\d+)\t(.+)$', p.text)
        if m and int(m.group(1)) <= 32:
            approved[int(m.group(1))] = m.group(2).strip().rstrip('.')
    assert len(approved) == 32
    approved_instr = next(p.text for p in v12.paragraphs if p.text.startswith('Please rate'))
    scale = '; '.join(f"{s['value']} – {s['label']}" for s in comps['sf6_1']['scale'])

    d = docx.Document()
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    d.add_heading('Sense Foraging Questionnaire — Study 1 (item pool 6)', level=1)
    d.add_paragraph(f'{AMENDMENT}. Protocol # 00051180. Replaces the 32-item Sense Foraging Scale (v1.2) for Study 1.')

    d.add_heading('Instructions shown before the statements', level=2)
    d.add_paragraph(DEF['instructions'])
    d.add_heading('Response scale', level=2)
    d.add_paragraph(scale + '. Every statement also offers a small "Prefer not to answer" option; the two '
                    'attention checks do not.')
    d.add_heading('Presentation', level=2)
    for line in [
        '48 statements and two attention checks (50 in all), shown about 8 per page on 6 pages, with every '
        'response option labelled.',
        'Each participant sees the statements in their own random order, drawn under fixed rules: statements '
        'from the same hypothesized factor are spread apart; statements worded in the opposite direction are '
        'spread through the questionnaire and never on the same page as the statement they mirror; one '
        'attention check falls in each half.',
        'Statements 35–48 are worded in the opposite direction (reverse-keyed) and are reverse-scored in '
        'analysis. They allow the analysis to separate agreement with the content from a general tendency to '
        'agree.',
    ]:
        d.add_paragraph(line, style='List Bullet')

    d.add_heading('Statements', level=2)
    for n in range(1, 49):
        cid = f'sf6_{n}'
        d.add_paragraph(f'{n}\t{comps[cid]["question"]}' + ('   (reverse-keyed)' if cid in reversed_ids else ''))
    d.add_heading('Attention checks (placed among the statements)', level=2)
    for i, cid in enumerate(('sf6_attn_disagree', 'sf6_attn_agree'), 1):
        d.add_paragraph(f'A{i}\t{comps[cid]["question"]}')
    d.add_paragraph('Failing both checks excludes a response from analysis; the participant is paid regardless.')

    d.add_page_break()
    d.add_heading('Changes from the approved questionnaire (v1.2)', level=2)
    d.add_paragraph(f'Approved instructions: "{approved_instr}" Study 1 instructions: as above (the definition '
                    'of the two modes, then "In general, how much do you agree with each statement?"). The '
                    'response scale is unchanged, with "Prefer not to answer" added.')
    d.add_paragraph('Statements 1–32 keep their approved numbers. Most rewordings use one pair of names for the two '
                    'modes, "the sensing mode" and "the doing mode", in place of "receptive, sensory mode" and '
                    '"goal-oriented mode", and shorten the statement. Statements 33–48 are new.')
    t = d.add_table(rows=1, cols=4)
    t.style = 'Table Grid'
    for cell, h in zip(t.rows[0].cells, ('#', 'Approved v1.2 wording', 'Study 1 wording', 'Change')):
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
    for n in range(1, 49):
        cid = f'sf6_{n}'
        new = comps[cid]['question']
        old = approved.get(n, '—')
        if n == 19:
            change = 'Replaced (asks about noticing pleasant / unpleasant / neutral before reacting)'
        elif n in (33, 34):
            change = 'New'
        elif n >= 35:
            change = 'New, reverse-keyed'
        elif old.rstrip('.') == new.rstrip('.'):
            change = 'Unchanged'
        else:
            change = 'Reworded'
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text = str(n), old, new, change
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    d.add_paragraph()
    d.add_paragraph('Hypothesized factors (for analysis; not shown to participants): '
                    + '; '.join(f'{f} ({", ".join(i[4:] for i in ids)})' for f, ids in DEF['factors'].items()) + '.')
    d.save(OUT / 'Sense_Foraging_Questionnaire_Study1.docx')
    print('wrote Sense_Foraging_Questionnaire_Study1.docx')


def main():
    OUT.mkdir(exist_ok=True)
    consent()
    debrief()
    background()
    questionnaire()


if __name__ == '__main__':
    main()
