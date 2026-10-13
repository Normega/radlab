"""Word documents for the Study 2 amendment (REB #00051180), October 2026.

Writes, beside the approved Rev2 package, a folder Rev4_Study2_Amendment:
  - Consent_Form_v5_Study2          (Rev2/Consent_Form_v3.docx with the Study 2 paragraphs; clean + highlighted)
  - Debrief_Form_Study2             (Rev2/Debrief_Form.docx with the Study 2 paragraphs; clean + highlighted)
  - Followup_Consent_Form_Study2    (new: the optional ~2-week retest)
  - Followup_Debrief_Form_Study2    (new; followed by the approved Mental Health Resources Sheet)
  - Study2_Measures                 (new: the session in order, every new measure's instructions, response
                                     scale and items as shown, sources, licences and web adaptations)

The consent and debrief edits are the ones the live Study 2 forms carry (build_study2_forms.py), applied
to the approved Word files paragraph by paragraph; the build fails if an approved paragraph is not found
verbatim. Measure wording is read from study2/*.json and study1/sf-short-20.json, the live definitions.

Usage:  python scripts/sense_foraging/build_study2_amendment_docs.py
"""
import copy
import json
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from build_study2_forms import CONSENT, DEBRIEF, RETEST_CONSENT, RETEST_DEBRIEF_INTRO  # noqa: E402

ETHICS = Path(r'I:\My Drive\Ethics\Sense Foraging Questionnaire Zindel')
REV2 = ETHICS / 'Rev2'
OUT = ETHICS / 'Rev4_Study2_Amendment'
AMENDMENT = 'Study 2 amendment, October 2026'
NOTE = f'{AMENDMENT}: changed text is highlighted. All other text is unchanged from the approved Rev2 version.'


def load(rel):
    return json.loads((HERE / rel).read_text(encoding='utf-8'))


def set_text(p, text, mark):
    """Replace a paragraph's text, keeping its first run's formatting."""
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    runs[0].bold = None
    if mark:
        runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW


def edit_doc(src, dst_stem, edits):
    """edits: [(old paragraph text, [new paragraphs])]. Writes a clean and a highlighted copy."""
    for mark in (False, True):
        d = docx.Document(src)
        # Rev2 marked its own changes (the Borealis consent split) in bold body text; carried over,
        # that bold would read as part of this amendment.
        for p in d.paragraphs:
            if len(p.text) > 80 and not p.text.startswith('['):
                for r in p.runs:
                    r.bold = None
        for old, new in edits:
            hits = [p for p in d.paragraphs if p.text.strip() == old.strip()]
            assert len(hits) == 1, f'{src.name}: approved paragraph not found once: {old[:70]}'
            p = hits[0]
            set_text(p, new[0], mark)
            for extra in new[1:]:
                el = copy.deepcopy(p._p)
                p._p.addnext(el)
                p = docx.text.paragraph.Paragraph(el, p._parent)
                set_text(p, extra, mark)
        if mark:
            note = d.paragraphs[0].insert_paragraph_before(NOTE)
            note.runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
            note.runs[0].font.bold = True
        name = f'{dst_stem}{"_changes_highlighted" if mark else ""}.docx'
        d.save(OUT / name)
        print('wrote', name)


def consent():
    edits = [(old, new) for _, old, new in CONSENT]
    d0 = docx.Document(REV2 / 'Consent_Form_v3.docx')
    revised = next(p.text for p in d0.paragraphs if p.text.startswith('REVISED'))
    edits.append((revised, [f'REVISED ({AMENDMENT}): Study 2 administers the background questions, the 20-statement '
                            'Sense Foraging short form, 24 candidate statements, and the validity, wellbeing and '
                            'stress measures (about 20 minutes). A random subsample is invited to an optional ~3-minute '
                            'follow-up about two weeks later, under its own consent form.']))
    edit_doc(REV2 / 'Consent_Form_v3.docx', 'Consent_Form_v5_Study2', edits)


def debrief():
    edits = [(old, new) for _, old, new in DEBRIEF]
    edit_doc(REV2 / 'Debrief_Form.docx', 'Debrief_Form_Study2', edits)


def new_doc(title):
    d = docx.Document()
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    for line in ('UNIVERSITY OF TORONTO', 'Department of Psychology', title,
                 'Development and Validation of the Sense Foraging Scale'):
        d.add_paragraph(line).runs[0].bold = True
    d.add_paragraph(f'NEW ({AMENDMENT}). Protocol # 00051180.').runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW
    return d


def blocks(d, items):
    for tag, text in items:
        if tag == 'h2':
            d.add_heading(text, level=2)
        else:
            d.add_paragraph(text)


def followup():
    d = new_doc('Consent to Participate in Research: Optional Follow-up')
    blocks(d, RETEST_CONSENT[1:])    # the first block is the on-screen title, already in the header
    d.save(OUT / 'Followup_Consent_Form_Study2.docx')
    print('wrote Followup_Consent_Form_Study2.docx')
    d = new_doc('Debriefing Statement: Optional Follow-up')
    blocks(d, RETEST_DEBRIEF_INTRO)
    d.add_paragraph('Followed by the approved Mental Health and Wellbeing Resources sheet (Rev2, unchanged), '
                    'exactly as in the Study 2 debrief.')
    d.save(OUT / 'Followup_Debrief_Form_Study2.docx')
    print('wrote Followup_Debrief_Form_Study2.docx')


# Session order, with where each piece's approval comes from.
SESSION = [
    ('Background questions', 'sf-background', 'Unchanged from Study 1 (approved Rev3).'),
    ('Sense Foraging Scale, short form (20 statements)', 'sf-short-20',
     'NEW. The 20 statements are a subset of the 48 approved for Study 1 (Rev3), selected from the Study 1 '
     'factor analysis, wording unchanged. One fixed order for everyone, 3 pages.'),
    ('Twenty-four candidate statements (unscored)', 'sf2-sfs-candidates',
     'NEW. 19 statements written for a future version of the scale (including 5 about seeking out sensory experience), '
     'and 5 from the approved Study 1 pool, wording unchanged (items 2, 6, 14, 25, 28); analysed, not scored. Three pages '
     'of 8, in a random order per participant: statements written for the same part never adjacent, reversed statements '
     'spread across the pages.'),
    ('— Validity block: the measures below in a random order per participant —', None, None),
    ('FFMQ-15', 'sf2-ffmq15', 'As approved (Rev2 Validity Scales Appendix).'),
    ('MAIA-2, with its attention check', 'sf2-maia2', 'As approved (Rev2 Validity Scales Appendix).'),
    ('BIDR-16', 'sf2-bidr16', 'As approved (Rev2 Validity Scales Appendix).'),
    ('PHQ-4', 'sf2-phq4', 'As approved (Rev2 Validity Scales Appendix).'),
    ('Life satisfaction (one item)', None, 'CHANGED FORMAT. The approved wording (Rev2 Validity Scales Appendix), answered on '
     'a six-point face scale (unhappy to delighted) instead of the appendix’s 1–7.'),
    ('IUS-12', 'sf2-ius12', 'NEW.'),
    ('SBI, Savoring the Moment subscale', 'sf2-sbi-moment', 'NEW.'),
    ('BFI-2 Open-Mindedness domain, plus one attention check', 'sf2-bfi2-om', 'NEW.'),
    ('Experiences Questionnaire, Decentering', 'sf2-eq-decentering', 'NEW.'),
    ('PSS-4', 'sf2-pss4', 'NEW.'),
    ('Past-week questions (3)', 'sf2-past-week', 'NEW. Written for this study.'),
    ('— End of the validity block —', None, None),
    ('Debrief', None, 'Study 2 debrief (this amendment).'),
]

SOURCES = {
    'sf2-ius12': ('Carleton, R. N., Norton, M. A. P. J., & Asmundson, G. J. G. (2007). Fearing the unknown: A short '
                  'version of the Intolerance of Uncertainty Scale. Journal of Anxiety Disorders, 21, 105–117. Item text, '
                  'anchors and order from the form circulated by the authors\' labs (University of Regina theses).',
                  'No licence or permission statement from the authors; ePROVIDE lists copyright as "No information". '
                  'Widely used in research without stated permission.'),
    'sf2-sbi-moment': ('Bryant, F. B. (2003). Savoring Beliefs Inventory (SBI). Journal of Mental Health, 12(2), 175–196; '
                       'items and instructions from Bryant & Veroff (2007), Savoring, Appendix B. Savoring the Moment '
                       'subscale: items 2, 5, 8, 11, 14, 17, 20, 23 (4 positively, 4 negatively keyed), in published order.',
                       'No licence statement; the 2003 article says copies are available from the author on request. '
                       'Permission requested from Prof. Bryant.'),
    'sf2-bfi2-om': ('Soto, C. J., & John, O. P. (2017). The next Big Five Inventory (BFI-2). Journal of Personality and '
                    'Social Psychology, 113(1), 117–143. The 12 Open-Mindedness items in their BFI-2 positions.',
                    '"BFI-2 items copyright 2015 by Oliver P. John and Christopher J. Soto. Permission is granted for '
                    'personal and research use of the BFI-2."'),
    'sf2-eq-decentering': ('Fresco, D. M., Moore, M. T., van Dulmen, M. H. M., Segal, Z. V., Ma, S. H., Teasdale, J. D., & '
                           'Williams, J. M. G. (2007). Initial psychometric properties of the Experiences Questionnaire. '
                           'Behavior Therapy, 38, 234–246. The 11-item decentering form and instructions as published by '
                           'the developers (Naragon-Gainey, McMahon, Park & Fresco, 2025, Appendix 1).',
                           'Developers: prior permission is not required for research use; cite Fresco et al. (2007).'),
    'sf2-pss4': ('Cohen, S., Kamarck, T., & Mermelstein, R. (1983). A global measure of perceived stress. Journal of '
                 'Health and Social Behavior, 24, 385–396. PSS-4 items and instructions from the Cohen lab form.',
                 'Distributed by Mapi Research Trust (ePROVIDE); electronic administration requires a completed licence '
                 'agreement. Request submitted / to be submitted before launch.'),
}

ADAPTATIONS = [
    ('All measures', 'Every item offers a small "Prefer not to answer" option (as approved for Study 1); the two '
                     'attention checks do not.'),
    ('IUS-12', '"Please circle the number" → "Please select the number".'),
    ('SBI', '"please circle the one number" → "please select the one number". Gender/age fields on the paper form '
            'omitted (collected in the background questions).'),
    ('BFI-2', '"Please write a number next to each statement" → "Please select a number for each statement"; the '
              'closing "Did you write a number…" check omitted. Only the 12 Open-Mindedness items are given. One '
              'attention check ("select Agree a little") added among them; it is not a BFI-2 item.'),
    ('EQ Decentering', '"best describe" → "best describes" (a slip in the source) and "All the Time" → "All the time"; '
                       'otherwise verbatim.'),
    ('PSS-4', '"please indicate your response by placing an \'X\' over the circle representing HOW OFTEN you felt" → '
              '"please indicate how often you felt".'),
]


def questionnaire_section(d, slug, rel):
    q = load(rel)
    comps = [c for p in q['pages'] for c in p['components']]
    d.add_paragraph(f'Title shown: "{q["name"]}". Instructions shown: "{q.get("instructions", "")}"')
    scales = {tuple(s['label'] for s in c['scale']) for c in comps if c['type'] == 'likert'}
    for sc in scales:
        d.add_paragraph('Response options: ' + ' / '.join(sc) + '.')
    rev = set(q.get('scoring', {}).get('reverse_items', []) or q.get('reverse_keyed', []))
    rev |= {c['id'] for c in q.get('candidates', []) if c['reverse_keyed']}
    s1 = {c['id']: c['study1_item'] for c in q.get('candidates', []) if 'study1_item' in c}
    for c in comps:
        text = f'{c["id"]}\t{c["question"]}'
        if c['id'] in rev:
            text += '   (reverse-scored)'
        if 'attn' in c['id']:
            text += '   (attention check)'
        if c['id'] in s1:
            text += f'   (Study 1 item {s1[c["id"]][4:]}, approved wording)'
        d.add_paragraph(text)
        if c['type'] == 'multiple_choice':
            d.add_paragraph('Options: ' + ' / '.join(o['label'] for o in c['options']) + ' / Prefer not to answer',
                            style='List Bullet')


def measures():
    d = docx.Document()
    st = d.styles['Normal']
    st.font.name = 'Calibri'
    st.font.size = Pt(11)
    d.add_heading('Study 2 measures', level=1)
    d.add_paragraph(f'{AMENDMENT}. Protocol # 00051180. Online, single session, about 20 minutes, Prolific '
                    '(general and contemplative-practitioner postings, as in Study 1).')

    d.add_heading('Sample', level=2)
    d.add_paragraph('CHANGED. About 1,000 completers (about 600 on the general posting and 400 on the contemplative '
                    'posting), up from the approved 450–500 for Study 2. The larger sample lets the candidate statements '
                    'be chosen in one random half of the sample and checked in the other, each half the size of Study 1. '
                    'Recruitment is otherwise as approved: two postings that differ only in how the study is described, '
                    'no screening on contemplative practice, adults 18+ fluent in English, and Study 1 participants '
                    'excluded. Payment is unchanged (USD 12 per hour, prorated).')
    d.add_paragraph('CHANGED. No country-of-residence filter. The application named the United States, United Kingdom, '
                    'Canada, Australia, Ireland and New Zealand; Study 2 is open to adults fluent in English in any '
                    'country. The debrief already lists crisis lines for those six countries and points participants '
                    'anywhere else to Find a Helpline (findahelpline.com), a free directory of crisis lines by country.')

    d.add_heading('The session, in order', level=2)
    t = d.add_table(rows=1, cols=2)
    t.style = 'Table Grid'
    for cell, h in zip(t.rows[0].cells, ('Step', 'Status')):
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
    for label, _, status in SESSION:
        row = t.add_row().cells
        row[0].text, row[1].text = label, status or ''
    d.add_paragraph()
    d.add_paragraph('The validity block runs in a random order for each participant, so that no measure is always '
                    'answered when participants are freshest or most tired. Two attention checks in all (MAIA-2, '
                    'approved; BFI-2, new). Failing both excludes from analysis; the participant is paid regardless.')

    d.add_heading('Optional follow-up (about 2 weeks later)', level=2)
    d.add_paragraph('About 200 randomly chosen Study 2 completers are invited through Prolific (a custom allowlist), '
                    'with a 7-day window, aiming for about 150 completions. The follow-up repeats the 20-statement '
                    'short form only (about 3 minutes, paid at the same rate) under its own consent and debrief. The '
                    'two sets of answers are linked by a random code before Prolific IDs are deleted.')

    d.add_heading('Sources and licences of the new measures', level=2)
    for slug, (src, lic) in SOURCES.items():
        label = next(lbl for lbl, s, _ in SESSION if s == slug)
        d.add_paragraph(label).runs[0].bold = True
        d.add_paragraph('Source: ' + src, style='List Bullet')
        d.add_paragraph('Licence: ' + lic, style='List Bullet')

    d.add_heading('Adaptations for web administration', level=2)
    for who, what in ADAPTATIONS:
        p = d.add_paragraph(style='List Bullet')
        p.add_run(who + ': ').bold = True
        p.add_run(what)

    d.add_page_break()
    d.add_heading('Items as shown', level=2)
    d.add_paragraph('Measures already approved (background questions, FFMQ-15, MAIA-2, BIDR-16, PHQ-4, and the life-'
                    'satisfaction wording) are not repeated here.')
    for label, slug, status in SESSION:
        if not slug or not status.startswith('NEW'):
            continue
        rel = 'study1/sf-short-20.json' if slug == 'sf-short-20' else f'study2/{slug}.json'
        d.add_heading(label, level=3)
        questionnaire_section(d, slug, rel)
    d.save(OUT / 'Study2_Measures.docx')
    print('wrote Study2_Measures.docx')


def main():
    OUT.mkdir(exist_ok=True)
    consent()
    debrief()
    followup()
    measures()


if __name__ == '__main__':
    main()
